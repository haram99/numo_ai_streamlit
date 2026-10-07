# -*- coding: utf-8 -*-
"""
رحلة الذكاء الاصطناعي — أكاديمية نمو (نسخة Streamlit)
- تسجيل اسم الطالبة ورقم الجوال (0500000000)
- ثلاثة مستويات: أسئلة، تدريب آلة (KNN)، اكتشاف التحيّز
- حفظ بيانات اللاعبات ونتائجهن في ملف Excel (players.xlsx)
- لوحة مشرف محمية بكلمة مرور لتنزيل الملف

التشغيل محليًا:  streamlit run app.py
"""
import hmac
import math
import os
import re
import threading
import uuid
from datetime import datetime
from pathlib import Path

import altair as alt
import pandas as pd
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image

BASE = Path(__file__).parent
LOGO = BASE / "logo.jpg"
DATA_FILE = Path(os.environ.get("PLAYERS_XLSX", BASE / "players.xlsx"))

AR_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "0123456789" * 2)

st.set_page_config(
    page_title="رحلة الذكاء الاصطناعي | أكاديمية نمو",
    page_icon=Image.open(LOGO) if LOGO.exists() else "🤖",
    layout="centered",
)

st.markdown(
    """
<style>
@import url('https://fonts.googleapis.com/css2?family=Tajawal:wght@400;500;700&display=swap');
.stApp, .stApp p, .stApp label, .stApp button, .stApp input, .stApp h1, .stApp h2,
.stApp h3, .stApp h4, .stApp li, .stApp span.stMarkdown { font-family: 'Tajawal', sans-serif; }
.stApp { direction: rtl; text-align: right; }
[data-testid="stMarkdownContainer"], [data-testid="stCaptionContainer"],
[data-testid="stAlert"], label, .stRadio { direction: rtl; text-align: right; }
.block-container { max-width: 760px; padding-top: 1.5rem; }
div.stButton > button, div[data-testid="stFormSubmitButton"] > button {
    width: 100%; border-radius: 12px; font-weight: 700; padding: .6rem 1rem; }
input[aria-label^="رقم الجوال"] { direction: ltr; text-align: center; letter-spacing: 2px; }
</style>
""",
    unsafe_allow_html=True,
)

# ============================== المحتوى ==============================
LEVEL1 = [
    ("ما الذي يميّز الذكاء الاصطناعي عن البرنامج التقليدي؟",
     ["يتعلّم من البيانات ويحسّن أداءه", "يعمل بدون كهرباء", "لا يحتاج إلى بيانات أبدًا"], 0,
     "الذكاء الاصطناعي يتعلّم الأنماط من البيانات بدل أن نكتب له كل قاعدة يدويًا."),
    ("أي مما يلي مثال على تطبيق للذكاء الاصطناعي؟",
     ["الآلة الحاسبة البسيطة", "التعرّف على الوجه لفتح الجوال", "مصباح كهربائي"], 1,
     "التعرّف على الوجه يتعلّم من صور كثيرة ليميّز الوجوه."),
    ("ما أهم مادة خام لتدريب نموذج ذكاء اصطناعي؟",
     ["البيانات", "الألوان", "الصوت العالي"], 0,
     "جودة البيانات وكميتها تحدّد جودة النموذج: «مدخلات سيئة = نتائج سيئة»."),
    ("عندما يقترح يوتيوب فيديو يعجبك، فهو يستخدم:",
     ["الحظ", "التعلّم من سلوكك وسلوك المستخدمين (أنظمة التوصية)", "اختيار موظف لكل شخص"], 1,
     "أنظمة التوصية تحلّل ما شاهدته لتتوقّع ما تحبّينه."),
    ("ما معنى «تدريب النموذج»؟",
     ["إعطاؤه أمثلة ليتعلّم منها", "إطفاؤه وتشغيله", "رسم شكل له"], 0,
     "ندرّب النموذج بأمثلة معنونة (مثلاً: هذه تفاحة، وهذه موزة) فيكتشف الفروق."),
]

LEVEL3 = [
    ("درّبوا نظامًا للتعرّف على الوجوه بصور أشخاص من بلد واحد فقط. ماذا سيحدث غالبًا؟",
     ["سيعمل بدقة متساوية مع الجميع", "سيخطئ أكثر مع وجوه لم يرها كثيرًا", "سيصبح أذكى من البشر"], 1,
     "هذا يسمّى «تحيّز البيانات»: النموذج يتقن ما رآه كثيرًا ويضعف فيما رآه قليلًا."),
    ("لتقليل التحيّز في النموذج، الأفضل أن:",
     ["نجمع بيانات متنوعة ومتوازنة", "نحذف نصف البيانات عشوائيًا", "نستخدم صورة واحدة فقط"], 0,
     "التنوّع والتوازن في البيانات يجعلان النموذج أعدل وأدق."),
    ("نظام يرفض طلبات التوظيف اعتمادًا على بيانات قديمة فيها تمييز. ما المشكلة؟",
     ["النموذج يتعلّم التمييز نفسه من البيانات", "الحاسوب بطيء", "الشاشة صغيرة"], 0,
     "النموذج يعكس ما في بياناته؛ لذلك يجب مراجعة البيانات وعدم الاعتماد الأعمى على النتائج."),
    ("هل يجب أن يراجع إنسان القرارات المهمة التي يتخذها الذكاء الاصطناعي؟",
     ["لا، الآلة لا تخطئ أبدًا", "نعم، للمراجعة والمسؤولية", "فقط في أيام العطل"], 1,
     "الإنسان مسؤول عن القرار النهائي، خاصة في الصحة والتعليم والعمل."),
]

# (الطول, الاستدارة, الصنف الحقيقي)
TEST_SET = [
    (4.5, 6.5, "تفاحة"), (7.5, 3.0, "موزة"), (3.0, 9.0, "تفاحة"),
    (8.5, 1.5, "موزة"), (6.0, 4.0, "موزة"), (5.0, 6.0, "تفاحة"),
]
FRUIT_COLORS = ["#e03b3b", "#f2c200"]  # تفاحة، موزة

LEVELS = [
    (1, "المستوى 1: ما هو الذكاء الاصطناعي؟", "quiz1"),
    (2, "المستوى 2: درّبي الآلة 🍎🍌", "train"),
    (3, "المستوى 3: اكتشفي التحيّز ⚖️", "quiz3"),
]

# ============================== التخزين في Excel ==============================
HEADERS = ["المعرّف", "تاريخ التسجيل", "الاسم", "رقم الجوال", "درجة المستوى 1",
           "نقاط المستوى 2", "درجة المستوى 3", "مجموع النقاط", "آخر تحديث"]
LEVEL_COL = {1: "درجة المستوى 1", 2: "نقاط المستوى 2", 3: "درجة المستوى 3"}


@st.cache_resource
def file_lock() -> threading.Lock:
    """قفل مشترك بين جلسات المستخدمات حتى لا يتعارض الحفظ المتزامن."""
    return threading.Lock()


def _open_workbook():
    if DATA_FILE.exists():
        return load_workbook(DATA_FILE)
    wb = Workbook()
    ws = wb.active
    ws.title = "اللاعبات"
    ws.sheet_view.rightToLeft = True
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1B6F86")
        c.alignment = Alignment(horizontal="center")
    for letter, w in zip("ABCDEFGHI", [14, 18, 26, 16, 16, 16, 16, 14, 18]):
        ws.column_dimensions[letter].width = w
    return wb


def save_player(pid: str, fields: dict) -> None:
    """إنشاء صف للاعبة أو تحديثه (حسب المعرّف) ثم الحفظ بشكل آمن."""
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    with file_lock():
        wb = _open_workbook()
        ws = wb.active
        row = next((r for r in range(2, ws.max_row + 1)
                    if ws.cell(r, 1).value == pid), None)
        if row is None:
            ws.append([pid, now, "", "", None, None, None, 0, now])
            row = ws.max_row
        for col_name, value in fields.items():
            cell = ws.cell(row, HEADERS.index(col_name) + 1, value)
            if col_name == "رقم الجوال":
                cell.number_format = "@"  # يحفظ الصفر في بداية الرقم
        ws.cell(row, HEADERS.index("آخر تحديث") + 1, now)
        tmp = DATA_FILE.with_suffix(".tmp")
        wb.save(tmp)
        tmp.replace(DATA_FILE)


def persist() -> None:
    s = st.session_state
    fields = {LEVEL_COL[n]: s.scores[n] for n in (1, 2, 3)}
    fields["مجموع النقاط"] = total()
    try:
        save_player(s.pid, fields)
    except Exception:  # لا نوقف اللعبة إذا فشل الحفظ
        s.save_error = True


# ============================== الحالة ==============================
def init_state():
    defaults = {
        "stage": "login", "name": "", "phone": "", "pid": "",
        "scores": {1: None, 2: None, 3: None},
        "qi": 0, "qc": 0, "picked": None,
        "train": [], "test": None, "save_error": False,
    }
    for k, v in defaults.items():
        st.session_state.setdefault(k, v)


def total() -> int:
    return sum(v or 0 for v in st.session_state.scores.values())


def go(stage: str):
    st.session_state.stage = stage
    st.rerun()


def start_level(stage: str):
    s = st.session_state
    s.qi, s.qc, s.picked, s.train, s.test = 0, 0, None, [], None
    go(stage)


def knn_predict(train, point, k=1):
    dists = sorted((math.dist((x, y), point), lbl) for x, y, lbl in train)
    top = [lbl for _, lbl in dists[:k]]
    return max(set(top), key=top.count)


# ============================== الشاشات ==============================
def header_bar():
    c1, c2 = st.columns([1, 5])
    if LOGO.exists():
        c1.image(str(LOGO), width=70)
    c2.markdown(f"**أكاديمية نمو | Numo Academy**  \n👩‍🎓 {st.session_state.name}")
    st.divider()


def screen_login():
    if LOGO.exists():
        _, mid, _ = st.columns([1, 1, 1])
        mid.image(str(LOGO), use_container_width=True)
    st.markdown("<h2 style='text-align:center'>مرحبًا بكِ في رحلة الذكاء الاصطناعي 🤖</h2>",
                unsafe_allow_html=True)
    with st.form("login"):
        name = st.text_input("اسم الطالبة")
        phone = st.text_input("رقم الجوال (مثال: 0500000000)", max_chars=10)
        consent = st.checkbox("أوافق على حفظ اسمي ورقم جوالي لدى أكاديمية نمو "
                              "لأغراض المتابعة والتواصل.")
        submitted = st.form_submit_button("ابدئي الرحلة ◀", type="primary")

    if not submitted:
        return
    name = " ".join(name.split())
    phone = phone.strip().translate(AR_DIGITS)
    if len(name) < 2 or not all(ch.isalpha() or ch == " " for ch in name):
        st.error("⚠️ فضلًا اكتبي اسمك بالحروف فقط (حرفان على الأقل).")
    elif not re.fullmatch(r"05\d{8}", phone):
        st.error("⚠️ رقم الجوال يجب أن يكون بالصيغة 0500000000 (10 أرقام تبدأ بـ 05).")
    elif not consent:
        st.error("⚠️ فضلًا وافقي على حفظ البيانات للمتابعة.")
    else:
        s = st.session_state
        s.name, s.phone, s.pid = name, phone, uuid.uuid4().hex[:10]
        try:
            save_player(s.pid, {"الاسم": name, "رقم الجوال": phone})
        except Exception:
            s.save_error = True
        go("home")


def screen_home():
    s = st.session_state
    st.markdown(f"### أهلًا بكِ يا {s.name} 👋")
    st.metric("مجموع نقاطك", total())
    st.write("ثلاثة مستويات لتتعلّمي أساسيات الذكاء الاصطناعي بالمرح!")
    for n, title, stage in LEVELS:
        done = s.scores[n] is not None
        if st.button(("✅ " if done else "") + title, key=f"lvl{n}"):
            start_level(stage)
    if all(v is not None for v in s.scores.values()):
        st.balloons()
        st.success(f"🏆 أنهيتِ كل المستويات يا {s.name}! مجموع نقاطك {total()}.")
    st.divider()
    if st.button("🚪 إنهاء الجلسة"):
        for k in list(st.session_state.keys()):
            del st.session_state[k]
        st.rerun()
    if s.save_error:
        st.warning("تعذّر حفظ النتيجة في الملف، أبلغي المشرفة.")


def screen_quiz(n: int):
    s = st.session_state
    qs = LEVEL1 if n == 1 else LEVEL3
    if s.qi >= len(qs):
        st.markdown(f"## 🎉 أحسنتِ يا {s.name}!")
        st.write(f"إجاباتك الصحيحة: **{s.qc} من {len(qs)}**  |  نقاط المستوى: **{s.scores[n]}**")
        if st.button("العودة للقائمة", type="primary"):
            go("home")
        return

    q, options, ans, explain = qs[s.qi]
    st.caption(f"المستوى {n} — سؤال {s.qi + 1} من {len(qs)}")
    st.markdown(f"#### {q}")
    for j, opt in enumerate(options):
        if st.button(opt, key=f"q{n}_{s.qi}_{j}", disabled=s.picked is not None):
            s.picked = j
            if j == ans:
                s.qc += 1
            st.rerun()

    if s.picked is not None:
        (st.success if s.picked == ans else st.error)(
            ("✅ إجابة صحيحة! " if s.picked == ans else "❌ ليست صحيحة. ") + explain)
        last = s.qi == len(qs) - 1
        if st.button("إنهاء المستوى ✔" if last else "التالي ◀", type="primary",
                     key=f"next{n}_{s.qi}"):
            s.qi += 1
            s.picked = None
            if s.qi >= len(qs):
                new = s.qc * 10
                s.scores[n] = max(s.scores[n] or 0, new)  # نحتفظ بأفضل نتيجة
                persist()
            st.rerun()
    if st.button("◀ القائمة", key=f"back{n}"):
        go("home")


def draw_chart():
    s = st.session_state
    if not s.train:
        st.info("أضيفي أمثلة لتظهر هنا على الرسم.")
        return
    color = alt.Color("النوع:N", legend=alt.Legend(title=None),
                      scale=alt.Scale(domain=["تفاحة", "موزة"], range=FRUIT_COLORS))
    x = alt.X("الطول:Q", scale=alt.Scale(domain=[0, 10]))
    y = alt.Y("الاستدارة:Q", scale=alt.Scale(domain=[0, 10]))
    df = pd.DataFrame(s.train, columns=["الطول", "الاستدارة", "النوع"])
    chart = alt.Chart(df).mark_circle(size=230, opacity=0.9, stroke="black",
                                      strokeWidth=1).encode(x=x, y=y, color=color)
    if s.test and "preds" in s.test:
        tdf = pd.DataFrame(s.test["preds"],
                           columns=["الطول", "الاستدارة", "النوع", "الحقيقي", "النتيجة"])
        squares = alt.Chart(tdf).mark_square(size=300, strokeWidth=4).encode(
            x=x, y=y, color=color,
            stroke=alt.Stroke("النتيجة:N", legend=alt.Legend(title="اختبار الآلة"),
                              scale=alt.Scale(domain=["صحيح", "خطأ"],
                                              range=["#2e9e5b", "#000000"])),
            tooltip=["الطول", "الاستدارة", "النوع", "الحقيقي", "النتيجة"])
        chart = chart + squares
    st.altair_chart(chart.properties(height=320), use_container_width=True)


def screen_train():
    s = st.session_state
    st.markdown("#### 🍎🍌 درّبي الآلة على التفريق بين التفاحة والموزة")
    st.write("حرّكي المنزلقين لوصف الفاكهة ثم اختاري اسمها. كلما أعطيتِ الآلة أمثلة أكثر "
             "وأكثر تنوّعًا، تعلّمت أفضل!")
    c1, c2 = st.columns(2)
    length = c1.slider("الطول (قصير ← طويل)", 0.0, 10.0, 5.0, 0.5)
    roundness = c2.slider("الاستدارة (مستقيم ← مستدير)", 0.0, 10.0, 5.0, 0.5)

    b1, b2, b3, b4 = st.columns(4)
    if b1.button("تفاحة 🍎"):
        s.train.append((length, roundness, "تفاحة"))
        s.test = None
    if b2.button("موزة 🍌"):
        s.train.append((length, roundness, "موزة"))
        s.test = None
    if b3.button("🗑️ مسح"):
        s.train, s.test = [], None
    if b4.button("اختبري ✅", type="primary"):
        labels = {lbl for *_, lbl in s.train}
        if len(labels) < 2 or len(s.train) < 4:
            s.test = {"error": "⚠️ أضيفي 4 أمثلة على الأقل، ومن النوعين معًا (تفاحة وموزة)."}
        else:
            preds = []
            for x, y, true in TEST_SET:
                p = knn_predict(s.train, (x, y))
                preds.append((x, y, p, true, "صحيح" if p == true else "خطأ"))
            right = sum(1 for *_, ok in preds if ok == "صحيح")
            s.test = {"preds": preds, "right": right,
                      "counts": {l: sum(1 for *_, t in s.train if t == l) for l in labels}}
            s.scores[2] = max(s.scores[2] or 0, right * 5)  # حتى 30 نقطة
            persist()

    draw_chart()
    st.caption(f"أمثلة التدريب: {len(s.train)}")

    t = s.test
    if t:
        if "error" in t:
            st.error(t["error"])
        else:
            n = len(TEST_SET)
            msg = (f"دقة الآلة: **{t['right']} من {n}** "
                   "(المربعات = اختبار جديد، الإطار الأخضر = تنبؤ صحيح)")
            if t["right"] == n:
                st.success("🏆 ممتاز! الآلة تعلّمت من أمثلتك. " + msg)
            else:
                st.warning(msg + "  \nجرّبي إضافة أمثلة أكثر وأكثر تنوّعًا قرب الحدود بين النوعين.")
            c = t["counts"]
            if max(c.values()) >= 3 * min(c.values()):
                st.info("⚖️ لاحظي: أمثلتك غير متوازنة بين النوعين، وهذا قد يسبّب تحيّزًا في النموذج!")
    if st.button("◀ القائمة", key="back2"):
        go("home")


# ============================== لوحة المشرف ==============================
def admin_password():
    try:
        if "ADMIN_PASSWORD" in st.secrets:
            return str(st.secrets["ADMIN_PASSWORD"])
    except Exception:
        pass
    return os.environ.get("ADMIN_PASSWORD")


def admin_panel():
    pw = admin_password()
    if not pw:
        return
    with st.sidebar.expander("🔒 لوحة المشرفة"):
        entered = st.text_input("كلمة المرور", type="password", key="admin_pw")
        if not entered:
            return
        if not hmac.compare_digest(entered.encode(), pw.encode()):
            st.error("كلمة المرور غير صحيحة.")
            return
        if not DATA_FILE.exists():
            st.info("لا توجد بيانات بعد.")
            return
        with file_lock():
            data = DATA_FILE.read_bytes()
        st.download_button("⬇️ تنزيل ملف اللاعبات (Excel)", data, file_name="players.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        st.dataframe(pd.read_excel(DATA_FILE, dtype={"رقم الجوال": str}),
                     use_container_width=True)


# ============================== التشغيل ==============================
def main():
    init_state()
    admin_panel()
    stage = st.session_state.stage
    if stage != "login":
        header_bar()
    if stage == "login":
        screen_login()
    elif stage == "home":
        screen_home()
    elif stage == "quiz1":
        screen_quiz(1)
    elif stage == "quiz3":
        screen_quiz(3)
    elif stage == "train":
        screen_train()


main()
